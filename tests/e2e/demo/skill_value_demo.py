"""Skill value demo: guided failure -> auto-generated skill -> unguided success.

Design goals (see docs/testing/demo-skill-generation-case.md and the
stability reports):

1. Before (no skill): give an underspecified real-business task so the agent
   naturally drifts (bash instead of doc tools, ad-hoc sheet/header, no
   read-back). A human then guides it with the business conventions.
2. Generate: the guided history is distilled into a SKILL.md.
3. After (skill injected): give a same-family but different task and require
   zero human intervention; the skill must carry the conventions.

Reports a comparison of human intervention rounds, tool calls, read-back
behaviour, latency and produced file checks.

Usage (demo stack running):
    .venv/Scripts/python.exe tests/e2e/demo/skill_value_demo.py
"""

from __future__ import annotations

import json
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx

PROJECT_ROOT = Path(__file__).resolve().parents[4]
BASE_URL = "http://127.0.0.1:8900"
WORKSPACE = Path(r"D:\workspace")
ROUND_TIMEOUT = 900
SUMMARY_TIMEOUT = 1500

BEFORE_TASK = (
    "你是后台运营专员。工作区里有租户用量原始数据文件 tenant-usage-raw.csv，"
    "请把它整理成一份 Excel 周报（输出文件 tenant-usage-weekly.xlsx）："
    "数据要准确、格式要规范，适合直接发给各租户。完成后 task_done 并回复文件路径。"
)

GUIDE_TASK = (
    "运营规范补充（重要，请按此重做）：\n"
    "1）必须使用 excel_edit_tool 文档工具创建 Excel，禁止用 bash 或脚本生成文件；\n"
    "2）Sheet 名固定为 W2026-35；\n"
    "3）表头固定为：租户 / 会话数 / 输入Token / 输出Token / 成功会话 / 成功率%，"
    "其中 成功率% = 成功会话 ÷ 会话数 × 100（保留 1 位小数）；\n"
    "4）创建后必须用 excel_edit_tool 的 read 回读核对表头、数据行数与最后一行内容。\n"
    "请重新生成 tenant-usage-weekly.xlsx，task_done 并回复文件路径。"
)

AFTER_TASK = (
    "你是后台运营专员。工作区里有新一周的租户用量原始数据 tenant-usage-raw-2.csv，"
    "请按运营规范整理成 Excel 周报（输出文件 tenant-usage-weekly-w36.xlsx），"
    "完成后 task_done 并回复文件路径。"
)


def _state(client: httpx.Client) -> dict[str, Any]:
    resp = client.get(f"{BASE_URL}/api/state")
    resp.raise_for_status()
    return resp.json()


def _wait_round(
    client: httpx.Client, before: int, timeout: int = ROUND_TIMEOUT
) -> dict[str, Any]:
    deadline = time.time() + timeout
    while time.time() < deadline:
        state = _state(client)
        rounds = state.get("rounds") or []
        if len(rounds) > before and rounds[-1].get("state") in ("completed", "failed"):
            return rounds[-1]
        time.sleep(3)
    raise TimeoutError("round did not complete")


def _wait_summary(client: httpx.Client, timeout: int = SUMMARY_TIMEOUT) -> dict[str, Any]:
    deadline = time.time() + timeout
    while time.time() < deadline:
        summary = _state(client).get("summary") or {}
        if summary.get("status") in ("completed", "failed"):
            return summary
        time.sleep(3)
    raise TimeoutError("summary did not finish")


def _check_workbook(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    if not path.is_file():
        return {"exists": False}
    workbook = load_workbook(str(path))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    return {
        "exists": True,
        "sheet": sheet.title,
        "header": list(rows[0]) if rows else [],
        "row_count": len(rows),
        "first_data": list(rows[1]) if len(rows) > 1 else [],
    }


def main() -> None:
    report: dict[str, Any] = {
        "started_at": datetime.now().isoformat(timespec="seconds"),
        "before": None,
        "generation": None,
        "after": None,
    }
    with httpx.Client(timeout=60) as client:
        # ---- Phase A: guided session without skill ----
        resp = client.post(f"{BASE_URL}/api/start", json={"auto_approve": True, "skills": []})
        resp.raise_for_status()

        before = len(_state(client)["rounds"])
        resp = client.post(f"{BASE_URL}/api/round", json={"task": BEFORE_TASK})
        resp.raise_for_status()
        t0 = time.time()
        r1 = _wait_round(client, before)
        report["before"] = {
            "task_kind": "underspecified (no skill)",
            "state": r1.get("state"),
            "seconds": round(time.time() - t0, 1),
            "tool_calls": r1.get("tool_calls") or [],
            "message_head": (r1.get("message") or "")[:220],
        }
        print(
            f"[skill-value] before round -> {r1.get('state')} "
            f"tools={sorted(set(r1.get('tool_calls') or []))}",
            flush=True,
        )

        before = len(_state(client)["rounds"])
        resp = client.post(f"{BASE_URL}/api/round", json={"task": GUIDE_TASK})
        resp.raise_for_status()
        t0 = time.time()
        r2 = _wait_round(client, before)
        report["guide"] = {
            "human_intervention_rounds": 1,
            "state": r2.get("state"),
            "seconds": round(time.time() - t0, 1),
            "tool_calls": r2.get("tool_calls") or [],
        }
        print(
            f"[skill-value] guide round -> {r2.get('state')} "
            f"tools={sorted(set(r2.get('tool_calls') or []))}",
            flush=True,
        )

        # ---- Generate skill from the guided history ----
        resp = client.post(f"{BASE_URL}/api/summarize")
        resp.raise_for_status()
        t0 = time.time()
        summary = _wait_summary(client)
        report["generation"] = {
            "status": summary.get("status"),
            "attempt": summary.get("attempt"),
            "skill_id": summary.get("skill_id"),
            "file_path": summary.get("file_path"),
            "seconds": round(time.time() - t0, 1),
        }
        if summary.get("status") != "completed":
            raise SystemExit(f"generation failed: {summary.get('error')}")
        skill_id = summary["skill_id"]
        skill_md = (PROJECT_ROOT / summary["file_path"]).read_text(encoding="utf-8")
        report["generation"]["skill_markdown"] = skill_md[:900]

        # ---- Phase B: unguided session with the generated skill ----
        resp = client.post(
            f"{BASE_URL}/api/start",
            json={"auto_approve": True, "skills": [skill_id]},
        )
        resp.raise_for_status()
        before = len(_state(client)["rounds"])
        resp = client.post(f"{BASE_URL}/api/round", json={"task": AFTER_TASK})
        resp.raise_for_status()
        t0 = time.time()
        r3 = _wait_round(client, before)
        report["after"] = {
            "task_kind": "same-family different task (skill injected)",
            "state": r3.get("state"),
            "human_intervention_rounds": 0,
            "seconds": round(time.time() - t0, 1),
            "tool_calls": r3.get("tool_calls") or [],
            "message_head": (r3.get("message") or "")[:220],
            "workbook": _check_workbook(WORKSPACE / "tenant-usage-weekly-w36.xlsx"),
        }
        print(
            f"[skill-value] after round -> {r3.get('state')} "
            f"tools={sorted(set(r3.get('tool_calls') or []))}",
            flush=True,
        )

    print(json.dumps(report, ensure_ascii=False, indent=2))
    out = PROJECT_ROOT / "docs" / "testing" / "stability-reports" / (
        datetime.now().strftime("%Y%m%d-%H%M%S") + "-skill-value.json"
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[skill-value] report written to {out}")


if __name__ == "__main__":
    main()
